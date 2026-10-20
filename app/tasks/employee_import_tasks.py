from __future__ import annotations

from datetime import UTC, datetime
from typing import Any

from openpyxl import load_workbook
from pymongo.errors import DuplicateKeyError, PyMongoError

from app.database.mongodb import employee_import_jobs_collection, employees_collection
from app.tasks.celery_app import celery_app


def _job_collection():
    return employee_import_jobs_collection


@celery_app.task(
    bind=True,
    name="employee_import.process",
    max_retries=3,
)
def import_employees_from_excel(
    self,
    job_id: str,
    tenant_id: str,
    file_path: str,
) -> dict[str, Any]:
    jobs = _job_collection()
    job = jobs.find_one({"_id": job_id})
    if job is None:
        raise ValueError(f"Import job {job_id} not found")

    jobs.update_one(
        {"_id": job_id},
        {
            "$set": {
                "status": "processing",
                "tenant_id": tenant_id,
                "started_at": datetime.now(UTC).isoformat(),
            }
        },
    )

    try:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise ValueError("Excel file is empty")

        headers = [str(value).strip().lower() for value in rows[0]]
        required_columns = {"employee_id", "name", "email"}
        missing = sorted(required_columns - set(headers))
        if missing:
            raise ValueError(f"Missing required columns: {', '.join(missing)}")

        total_rows = len(rows) - 1
        processed = 0
        succeeded = 0
        failed = 0
        errors: list[str] = []
        jobs.update_one(
            {"_id": job_id},
            {"$set": {"total_rows": total_rows}},
        )

        for row in rows[1:]:
            processed += 1
            record = dict(zip(headers, row))
            if not any(value not in (None, "") for value in record.values()):
                continue

            employee_id = str(record.get("employee_id") or "").strip()
            name = str(record.get("name") or "").strip()
            email = str(record.get("email") or "").strip().lower()
            department = str(record.get("department") or "Unassigned").strip()

            if not employee_id or not name or not email:
                failed += 1
                errors.append(
                    f"Row {processed}: employee_id, name and email are required"
                )
                continue

            try:
                employees_collection.insert_one(
                    {
                        "tenant_id": tenant_id,
                        "employee_id": employee_id,
                        "name": name,
                        "email": email,
                        "department": department,
                        "is_active": True,
                    }
                )
                succeeded += 1
            except DuplicateKeyError:
                failed += 1
                errors.append(
                    f"Row {processed}: duplicate employee or email for tenant"
                )

        final_status = "completed_with_errors" if errors else "completed"
        jobs.update_one(
            {"_id": job_id},
            {
                "$set": {
                    "status": final_status,
                    "processed": processed,
                    "succeeded": succeeded,
                    "failed": failed,
                    "errors": errors[:100],
                    "completed_at": datetime.now(UTC).isoformat(),
                }
            },
        )
        return {
            "job_id": job_id,
            "status": final_status,
            "processed": processed,
            "succeeded": succeeded,
            "failed": failed,
        }
    except (OSError, ValueError) as exc:
        jobs.update_one(
            {"_id": job_id},
            {
                "$set": {
                    "status": "failed",
                    "error": str(exc),
                    "completed_at": datetime.now(UTC).isoformat(),
                }
            },
        )
        raise
    except PyMongoError as exc:
        jobs.update_one(
            {"_id": job_id},
            {"$set": {"status": "retrying", "last_error": type(exc).__name__}},
        )
        raise self.retry(exc=exc, countdown=10)
